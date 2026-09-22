# -*- coding: utf-8 -*-
"""
StudyAbroad.AI Document Exporter
Converts StudyAbroad_AI_Project_Plan.md and StudyAbroad_AI_SRS.md
into client-ready PDF and Word (.docx) documents.
"""
import os
import re
import subprocess
import markdown
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOCS_DIR = r"E:\Project\Ai Projects\New folder (2)\docs"
EDGE_PATH = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"

DOCUMENTS = [
    {
        "md": os.path.join(DOCS_DIR, "StudyAbroad_AI_Project_Plan.md"),
        "pdf": os.path.join(DOCS_DIR, "StudyAbroad_AI_Project_Plan.pdf"),
        "docx": os.path.join(DOCS_DIR, "StudyAbroad_AI_Project_Plan.docx"),
        "title": "StudyAbroad.AI — Master Project Plan & Strategy",
        "subtitle": "Autonomous Multi-Agent Global Education Platform (12 Autonomous AI Agents)",
        "badge": "CONFIDENTIAL & PROPRIETARY CLIENT BRIEFING"
    },
    {
        "md": os.path.join(DOCS_DIR, "StudyAbroad_AI_SRS.md"),
        "pdf": os.path.join(DOCS_DIR, "StudyAbroad_AI_SRS.pdf"),
        "docx": os.path.join(DOCS_DIR, "StudyAbroad_AI_SRS.docx"),
        "title": "StudyAbroad.AI — Software Requirements Specification (SRS)",
        "subtitle": "IEEE 830 Architecture, Data Models, Multi-Agent Orchestration & Security",
        "badge": "ENTERPRISE TECHNICAL SPECIFICATION (IEEE 830)"
    }
]

CSS_STYLES = """
@import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@300;400;500;600;700;800&family=JetBrains+Mono:wght@400;500&display=swap');

@page {
    size: A4;
    margin: 18mm 18mm 18mm 18mm;
    @bottom-right {
        content: counter(page);
    }
}

* {
    box-sizing: border-box;
    -webkit-print-color-adjust: exact !important;
    print-color-adjust: exact !important;
}

body {
    font-family: 'Plus Jakarta Sans', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif;
    color: #1e293b;
    background: #ffffff;
    line-height: 1.65;
    font-size: 13px;
    margin: 0;
    padding: 0;
}

.doc-cover {
    background: linear-gradient(135deg, #0f172a 0%, #1e1b4b 50%, #312e81 100%);
    color: #ffffff;
    padding: 34px 30px;
    border-radius: 12px;
    margin-bottom: 26px;
    box-shadow: 0 10px 25px -5px rgba(15, 23, 42, 0.2);
}

.doc-badge {
    display: inline-block;
    background: rgba(99, 102, 241, 0.25);
    border: 1px solid #818cf8;
    color: #c7d2fe;
    padding: 4px 12px;
    border-radius: 9999px;
    font-size: 10.5px;
    font-weight: 700;
    letter-spacing: 0.08em;
    text-transform: uppercase;
    margin-bottom: 12px;
}

.doc-cover h1 {
    color: #ffffff !important;
    font-size: 23px;
    font-weight: 800;
    margin: 0 0 8px 0;
    border-bottom: none !important;
    padding: 0;
    letter-spacing: -0.02em;
}

.doc-cover p {
    color: #cbd5e1;
    font-size: 13px;
    margin: 0;
    font-weight: 400;
}

.doc-meta {
    display: flex;
    gap: 20px;
    margin-top: 16px;
    padding-top: 12px;
    border-top: 1px solid rgba(255,255,255,0.15);
    font-size: 11px;
    color: #94a3b8;
}

.doc-meta strong {
    color: #e2e8f0;
}

h1 {
    color: #0f172a;
    font-size: 18px;
    font-weight: 700;
    margin-top: 28px;
    margin-bottom: 10px;
    padding-bottom: 5px;
    border-bottom: 2px solid #e2e8f0;
    page-break-after: avoid;
}

h2 {
    color: #1e3a8a;
    font-size: 15px;
    font-weight: 700;
    margin-top: 20px;
    margin-bottom: 8px;
    page-break-after: avoid;
}

h3 {
    color: #334155;
    font-size: 13px;
    font-weight: 600;
    margin-top: 15px;
    margin-bottom: 6px;
    page-break-after: avoid;
}

h4 {
    color: #475569;
    font-size: 12px;
    font-weight: 600;
    margin-top: 12px;
    margin-bottom: 4px;
}

p {
    margin: 0 0 9px 0;
    color: #334155;
}

ul, ol {
    margin: 0 0 10px 0;
    padding-left: 20px;
    color: #334155;
}

li {
    margin-bottom: 3.5px;
}

strong {
    color: #0f172a;
    font-weight: 600;
}

code {
    font-family: 'JetBrains Mono', Consolas, Monaco, monospace;
    background-color: #f1f5f9;
    color: #0f172a;
    padding: 2px 5px;
    border-radius: 4px;
    font-size: 11.5px;
    border: 1px solid #e2e8f0;
}

pre {
    background-color: #0f172a;
    color: #f8fafc;
    padding: 12px 14px;
    border-radius: 8px;
    overflow-x: auto;
    font-size: 11.5px;
    line-height: 1.5;
    margin: 12px 0;
    page-break-inside: avoid;
    border: 1px solid #1e293b;
}

pre code {
    background: transparent;
    color: inherit;
    padding: 0;
    border: none;
}

table {
    width: 100%;
    border-collapse: separate;
    border-spacing: 0;
    margin: 14px 0;
    font-size: 11.5px;
    page-break-inside: avoid;
    border: 1px solid #cbd5e1;
    border-radius: 7px;
    overflow: hidden;
}

th {
    background-color: #1e293b;
    color: #ffffff;
    font-weight: 600;
    text-align: left;
    padding: 8px 10px;
    border-bottom: 1px solid #0f172a;
    font-size: 11.5px;
}

td {
    padding: 7px 10px;
    border-bottom: 1px solid #e2e8f0;
    color: #334155;
    vertical-align: top;
}

tr:last-child td {
    border-bottom: none;
}

tr:nth-child(even) td {
    background-color: #f8fafc;
}

blockquote {
    border-left: 4px solid #6366f1;
    background-color: #eef2ff;
    color: #312e81;
    margin: 12px 0;
    padding: 8px 14px;
    border-radius: 0 7px 7px 0;
    font-size: 12px;
    page-break-inside: avoid;
}

blockquote p:last-child {
    margin-bottom: 0;
}

hr {
    border: 0;
    height: 1px;
    background: #e2e8f0;
    margin: 20px 0;
}

.footer-note {
    margin-top: 32px;
    padding-top: 10px;
    border-top: 1px solid #e2e8f0;
    text-align: center;
    font-size: 10px;
    color: #94a3b8;
    page-break-inside: avoid;
}
"""

def generate_pdf(item):
    md_path = item["md"]
    pdf_path = item["pdf"]
    title = item["title"]
    subtitle = item["subtitle"]
    badge = item["badge"]

    with open(md_path, "r", encoding="utf-8") as f:
        md_content = f.read()

    html_body = markdown.markdown(
        md_content,
        extensions=["tables", "fenced_code", "nl2br", "sane_lists"]
    )

    full_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>{title}</title>
    <style>{CSS_STYLES}</style>
</head>
<body>
    <div class="doc-cover">
        <div class="doc-badge">{badge}</div>
        <h1>{title}</h1>
        <p>{subtitle}</p>
        <div class="doc-meta">
            <span><strong>System:</strong> StudyAbroad.AI Platform</span>
            <span><strong>Architecture:</strong> 12 Autonomous Agents</span>
            <span><strong>Distribution:</strong> Confidential Client Copy</span>
        </div>
    </div>
    
    <div class="doc-content">
        {html_body}
    </div>
    
    <div class="footer-note">
        StudyAbroad.AI — Autonomous Multi-Agent Global Education Intelligence Ecosystem • Confidential Client Document
    </div>
</body>
</html>"""

    html_path = md_path.replace(".md", ".html")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(full_html)

    cmd = [
        EDGE_PATH,
        "--headless",
        "--disable-gpu",
        "--allow-file-access-from-files",
        "--run-all-compositor-stages-before-draw",
        "--no-pdf-header-footer",
        f"--print-to-pdf={pdf_path}",
        f"file:///{html_path.replace(os.sep, '/')}"
    ]
    subprocess.run(cmd, capture_output=True)

    if os.path.exists(pdf_path):
        size_kb = os.path.getsize(pdf_path) / 1024
        print(f"  [OK] PDF created:  {os.path.basename(pdf_path)} ({size_kb:.1f} KB)")
    else:
        print(f"  [FAIL] PDF failed: {pdf_path}")

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def generate_docx(item):
    md_path = item["md"]
    docx_path = item["docx"]
    title = item["title"]
    subtitle = item["subtitle"]
    badge = item["badge"]

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # Executive Cover Table
    banner_table = doc.add_table(rows=1, cols=1)
    banner_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    banner_table.autofit = False

    cell = banner_table.cell(0, 0)
    set_cell_background(cell, "0F172A")
    set_cell_margins(cell, top=220, bottom=220, left=240, right=240)
    cell.width = Inches(6.9)

    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(3)
    run_badge = p0.add_run(badge)
    run_badge.font.name = 'Segoe UI'
    run_badge.font.size = Pt(8.5)
    run_badge.font.bold = True
    run_badge.font.color.rgb = RGBColor(0x81, 0x8C, 0xF8)

    p1 = cell.add_paragraph()
    p1.paragraph_format.space_before = Pt(2)
    p1.paragraph_format.space_after = Pt(3)
    run_title = p1.add_run(title)
    run_title.font.name = 'Segoe UI'
    run_title.font.size = Pt(15)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(8)
    run_sub = p2.add_run(subtitle)
    run_sub.font.name = 'Segoe UI'
    run_sub.font.size = Pt(9.5)
    run_sub.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    i = 0
    in_code_block = False
    code_lines = []

    while i < len(lines):
        line = lines[i].rstrip('\r\n')

        # Code block
        if line.strip().startswith('```'):
            if in_code_block:
                in_code_block = False
                code_text = '\n'.join(code_lines)
                code_lines = []

                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                c = tbl.cell(0, 0)
                set_cell_background(c, "1E293B")
                set_cell_margins(c, top=100, bottom=100, left=140, right=140)
                c.width = Inches(6.9)

                cp = c.paragraphs[0]
                cp.paragraph_format.space_before = Pt(0)
                cp.paragraph_format.space_after = Pt(0)
                crun = cp.add_run(code_text)
                crun.font.name = 'Consolas'
                crun.font.size = Pt(8.5)
                crun.font.color.rgb = RGBColor(0xF8, 0xFA, 0xFC)

                doc.add_paragraph().paragraph_format.space_after = Pt(3)
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Tables
        if line.strip().startswith('|') and '|' in line.strip()[1:]:
            table_raw = [line]
            i += 1
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_raw.append(lines[i].rstrip('\r\n'))
                i += 1

            parsed_rows = []
            for tr in table_raw:
                parts = [p.strip() for p in tr.strip().split('|')[1:-1]]
                if all(re.match(r'^:?-+:?$', p) for p in parts if p):
                    continue
                parsed_rows.append(parts)

            if parsed_rows:
                ncols = max(len(r) for r in parsed_rows)
                tbl = doc.add_table(rows=len(parsed_rows), cols=ncols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl.autofit = False

                for r_idx, row_data in enumerate(parsed_rows):
                    is_header = (r_idx == 0)
                    for c_idx in range(ncols):
                        c = tbl.cell(r_idx, c_idx)
                        text = row_data[c_idx] if c_idx < len(row_data) else ""

                        set_cell_margins(c, top=60, bottom=60, left=80, right=80)

                        if is_header:
                            set_cell_background(c, "1E293B")
                        elif r_idx % 2 == 1:
                            set_cell_background(c, "F8FAFC")
                        else:
                            set_cell_background(c, "FFFFFF")

                        p = c.paragraphs[0]
                        p.paragraph_format.space_before = Pt(1)
                        p.paragraph_format.space_after = Pt(1)

                        run = p.add_run(text)
                        run.font.name = 'Segoe UI'
                        run.font.size = Pt(8.5)
                        if is_header:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        else:
                            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

                doc.add_paragraph().paragraph_format.space_after = Pt(3)
            continue

        # Headers
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[2:].strip())
            run.font.name = 'Segoe UI'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        elif line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(11)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[3:].strip())
            run.font.name = 'Segoe UI'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        elif line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[4:].strip())
            run.font.name = 'Segoe UI'
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        elif line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[5:].strip())
            run.font.name = 'Segoe UI'
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        elif line.startswith('> '):
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            c = tbl.cell(0, 0)
            set_cell_background(c, "EEF2FF")
            set_cell_margins(c, top=70, bottom=70, left=120, right=120)
            c.width = Inches(6.9)
            cp = c.paragraphs[0]
            cp.paragraph_format.space_before = Pt(0)
            cp.paragraph_format.space_after = Pt(0)
            run = cp.add_run(line[2:].strip())
            run.font.name = 'Segoe UI'
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x31, 0x2E, 0x81)
            run.font.size = Pt(9.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif line.strip().startswith('- ') or line.strip().startswith('* '):
            text = line.strip()[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            for part in re.split(r'(\*\*.*?\*\*)', text):
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.name = 'Segoe UI'
                    r.font.bold = True
                    r.font.size = Pt(9.5)
                else:
                    r = p.add_run(part)
                    r.font.name = 'Segoe UI'
                    r.font.size = Pt(9.5)
        elif re.match(r'^\d+\.\s', line.strip()):
            match = re.match(r'^\d+\.\s', line.strip())
            text = line.strip()[len(match.group(0)):].strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            for part in re.split(r'(\*\*.*?\*\*)', text):
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.name = 'Segoe UI'
                    r.font.bold = True
                    r.font.size = Pt(9.5)
                else:
                    r = p.add_run(part)
                    r.font.name = 'Segoe UI'
                    r.font.size = Pt(9.5)
        elif line.strip() == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run("_______________________________________________________________________________")
            r.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
            r.font.size = Pt(7)
        elif line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2.5)
            for part in re.split(r'(\*\*.*?\*\*)', line):
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.name = 'Segoe UI'
                    r.font.bold = True
                    r.font.size = Pt(9.5)
                else:
                    r = p.add_run(part)
                    r.font.name = 'Segoe UI'
                    r.font.size = Pt(9.5)

        i += 1

    doc.save(docx_path)
    if os.path.exists(docx_path):
        size_kb = os.path.getsize(docx_path) / 1024
        print(f"  [OK] DOCX created: {os.path.basename(docx_path)} ({size_kb:.1f} KB)")
    else:
        print(f"  [FAIL] DOCX failed: {docx_path}")

def main():
    print("==================================================")
    print("  StudyAbroad.AI -- Document Exporter (PDF and DOCX)")
    print("==================================================")
    for item in DOCUMENTS:
        print(f"\nProcessing: {os.path.basename(item['md'])}...")
        generate_pdf(item)
        generate_docx(item)
    print("\n[COMPLETE] All documents generated successfully!")

if __name__ == "__main__":
    main()
